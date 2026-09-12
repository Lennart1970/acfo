"""Exact Excel sample writers. No pytest import — used by tests and pack scripts."""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

NL_HEADERS = [
    "Datum",
    "Boekstuk",
    "Leverancier",
    "Bedrag",
    "Valuta",
    "Btw-code",
    "Grootboek",
    "Journaal",
    "Betalingsconditie",
    "Omschrijving",
    "Administratie",
]

EN_HEADERS = [
    "EntryDate",
    "EntryNumber",
    "SupplierName",
    "AmountDC",
    "Currency",
    "VATCode",
    "GLAccountCode",
    "JournalCode",
    "PaymentCondition",
    "Description",
    "Division",
]


def monday_of(iso_year: int, iso_week: int) -> date:
    return date.fromisocalendar(iso_year, iso_week, 1)


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], header_row: int = 1):
    ws = wb.active if wb.active.title == "Sheet" and not wb.active["A1"].value else wb.create_sheet(title)
    if wb.active.title == "Sheet" and not wb.active["A1"].value:
        ws.title = title
    if header_row > 1:
        ws.cell(1, 1, "Exact Online — inkoop")
    for col, name in enumerate(headers, start=1):
        ws.cell(header_row, col, name)
    start = header_row + 1
    for r_i, row in enumerate(rows, start=start):
        for c_i, value in enumerate(row, start=1):
            ws.cell(r_i, c_i, value)
    return ws


def write_nl_export(path: Path) -> Path:
    """Multi-week Exact NL export used by tests and the sample fixture."""
    w36 = monday_of(2026, 36)
    w37 = monday_of(2026, 37)
    w38 = monday_of(2026, 38)
    rows = [
        [w36, "26400550", "DoiT International NL B.V.", 512.40, "EUR", "21", "4400", "60", "30 dagen", "Cloud september", "3919124"],
        [w36 + timedelta(days=1), "26400551", "DoiT International NL B.V.", 498.10, "EUR", "21", "4400", "60", "30 dagen", "Cloud usage", "3919124"],
        [w37, "26400556", "DoiT International NL B.V.", 505.00, "EUR", "21", "4400", "60", "30 dagen", "Cloud week 37", "3919124"],
        [w36 + timedelta(days=2), "26400552", "Office Supplies B.V.", 84.50, "EUR", "21", "4300", "60", "14 dagen", "Kantoor", "3919124"],
        [w36 + timedelta(days=3), "26400553", "Office Supplies B.V.", 91.20, "EUR", "21", "4300", "60", "14 dagen", "Kantoor", "3919124"],
        [w37 + timedelta(days=2), "26400557", "Office Supplies B.V.", 890.00, "EUR", "21", "4300", "60", "14 dagen", "Kantoor spoed", "3919124"],
        [w37 + timedelta(days=3), "26400558", "Nieuwe Crediteur B.V.", 1250.00, "EUR", "21", "4700", "60", "14 dagen", "Eenmalig advies", "3919124"],
        [w37 + timedelta(days=4), "26400559", "", 40.00, "EUR", "0", "4800", "60", "", "", "3919124"],
        [w38, "26400560", "DoiT International NL B.V.", 510.00, "EUR", "21", "4400", "60", "30 dagen", "Cloud week 38", "3919124"],
    ]
    wb = Workbook()
    add_sheet(wb, "Inkoop", NL_HEADERS, rows, header_row=2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_en_lines_export(path: Path) -> Path:
    day = date(2026, 9, 8)
    headers = EN_HEADERS + ["LineNumber"]
    rows = [
        [day, "1001", "Acme B.V.", 100.00, "EUR", "21", "4400", "60", "30 dagen", "Hosting", "4011167", 1],
        [day, "1001", "Acme B.V.", 20.00, "EUR", "21", "4500", "60", "30 dagen", "Support", "4011167", 2],
        [date(2026, 8, 11), "0990", "Acme B.V.", 110.00, "EUR", "21", "4400", "60", "30 dagen", "Hosting", "4011167", 1],
    ]
    wb = Workbook()
    add_sheet(wb, "TransactionLines", headers, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


INVANTIVE_HEADERS = [
    "Datum",
    "Dagboekcode",
    "Dagboekomschrijving",
    "Boekingnummer",
    "Regelnummer",
    "Bedrag Administratie Munteenheid",
    "Valuta",
    "Omschrijving",
    "Grootboekrekening Code",
    "Accountnaam",
    "BTW-code",
    "Grootboekrekeningsoort",
    "Divisie",
]


def write_invantive_purchases(path: Path) -> Path:
    """Invantive TransactionLines: double-entry purchase journals 40."""
    rows = [
        [date(2026, 5, 25), "40", "Purchases", 26400300, 1, 800.00, "EUR", "Lease mei", "4500", "Broekhuis Lease", "1", 120, 3919124],
        [date(2026, 5, 25), "40", "Purchases", 26400300, 9999, 168.00, "EUR", "Lease mei", "1450", "Broekhuis Lease", "1", 24, 3919124],
        [date(2026, 5, 25), "40", "Purchases", 26400300, 0, -968.00, "EUR", "Lease mei", "1300", "Broekhuis Lease", None, 22, 3919124],
        [date(2026, 6, 1), "40", "Purchases", 26400387, 1, 883.68, "EUR", "Broekhuis Lease", "4500", "Broekhuis Lease", "1", 120, 3919124],
        [date(2026, 6, 1), "40", "Purchases", 26400387, 9999, 185.57, "EUR", "Broekhuis Lease", "1450", "Broekhuis Lease", "1", 24, 3919124],
        [date(2026, 6, 1), "40", "Purchases", 26400387, 0, -1069.25, "EUR", "Broekhuis Lease", "1300", "Broekhuis Lease", None, 22, 3919124],
        [date(2026, 6, 1), "20", "NL06 RABO", 26200210, 1, 501.20, "EUR", "Reimbursement", "4380", "", None, 120, 3919124],
        [date(2026, 6, 1), "20", "NL06 RABO", 26200210, 0, -501.20, "EUR", "Reimbursement", "1100", "", None, 12, 3919124],
    ]
    wb = Workbook()
    add_sheet(wb, "TransactionLines", INVANTIVE_HEADERS, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_dutch_amounts_export(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Export"
    for col, name in enumerate(["Datum", "Boekstuk", "Leverancier", "Bedrag"], start=1):
        ws.cell(1, col, name)
    ws.cell(2, 1, "08-09-2026")
    ws.cell(2, 2, "77")
    ws.cell(2, 3, "Bakker B.V.")
    ws.cell(2, 4, "1.234,56")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
