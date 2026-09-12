from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from confidence_core import score_case
from excel_transactions import (
    apply_scope,
    build_cases,
    entries_in_week,
    list_weeks,
    load_entries,
    parse_number,
    parse_week,
)
from review_week import main, review_week, weeks_payload


def test_parse_week_variants():
    meta = parse_week("2026-W37")
    assert meta["week"] == "2026-W37"
    assert meta["start"] == date(2026, 9, 7)
    assert meta["end"] == date(2026, 9, 13)
    assert parse_week("week 37", today=date(2026, 9, 12))["week"] == "2026-W37"
    assert parse_week("2026-09-12")["week"] == "2026-W37"
    assert parse_week("37-2026")["week"] == "2026-W37"


def test_parse_dutch_thousands():
    assert parse_number("1.234,56") == 1234.56
    assert parse_number("€ 40,00") == 40.0


def test_load_nl_export_and_list_weeks(nl_export: Path):
    bundle = load_entries(nl_export)
    assert bundle["administrationCode"] == "3919124"
    assert len(bundle["entries"]) == 9
    weeks = {w["week"]: w["entries"] for w in list_weeks(bundle)}
    assert weeks["2026-W36"] == 4
    assert weeks["2026-W37"] == 4
    assert weeks["2026-W38"] == 1


def test_ask_week_without_flag(nl_export: Path, capsys):
    code = main(["--input", str(nl_export), "--json"])
    assert code == 3
    payload = json.loads(capsys.readouterr().out)
    assert payload["needs_week"] is True
    assert payload["ask"].startswith("Welke week")
    assert [w["week"] for w in payload["weeks"]] == ["2026-W36", "2026-W37", "2026-W38"]


def test_line_export_groups_and_sums(en_lines_export: Path):
    bundle = load_entries(en_lines_export)
    by_no = {e["entry"]["entryNumber"]: e for e in bundle["entries"]}
    assert len(by_no["1001"]["lines"]) == 2
    assert by_no["1001"]["entry"]["amountDC"] == 120.0
    assert bundle["administrationCode"] == "4011167"


def test_dutch_amount_cell(dutch_amounts_export: Path):
    bundle = load_entries(dutch_amounts_export)
    assert bundle["entries"][0]["entry"]["amountDC"] == 1234.56
    assert bundle["entries"][0]["entry"]["entryDate"] == "2026-09-08"


def test_week_37_routes_and_outcomes(nl_export: Path, tmp_path: Path):
    bundle = load_entries(nl_export)
    summary = review_week(bundle, "week 37")
    assert summary["week"] == "2026-W37"
    assert summary["entries"] == 4
    by_no = {r["entryNumber"]: r for r in summary["results"]}
    assert by_no["26400556"]["route"] == "Auto"
    assert by_no["26400556"]["confidence"] >= 0.9
    assert by_no["26400558"]["route"] == "Human Review"
    assert by_no["26400559"]["route"] == "Human Review"
    assert by_no["26400557"]["route"] in {"Human Review", "AI Review"}
    assert "vertrouwde boeking" in by_no["26400556"]["outcome"]
    assert "handmatig beoordelen" in by_no["26400558"]["outcome"]

    out = tmp_path / "Weekoverzicht-2026-W37.xlsx"
    code = main(["--input", str(nl_export), "--week", "2026-W37", "--output", str(out), "--json"])
    assert code == 0
    assert out.is_file()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Overzicht", "Boekingen", "Uitkomsten"]
    assert wb["Overzicht"]["B4"].value.startswith("week 37")
    assert wb["Overzicht"]["B10"].value == "nee"
    routes = [ws.value for ws in wb["Boekingen"]["K"][1:]]
    assert "Auto" in routes
    assert "Human Review" in routes
    assert wb["Uitkomsten"]["A4"].value is not None


def test_history_comes_from_other_weeks(nl_export: Path):
    bundle = load_entries(nl_export)
    _, matched = entries_in_week(bundle, "2026-W37")
    doit = next(e for e in matched if e["entry"]["entryNumber"] == "26400556")
    case = next(c for c in build_cases(bundle, [doit]) if c["entry"]["entryNumber"] == "26400556")
    hist_nos = {h["entryNumber"] for h in case["history"]}
    assert "26400550" in hist_nos
    assert "26400551" in hist_nos
    assert "26400556" not in hist_nos
    scored = score_case(case)
    assert scored["signals"]["booked_before"] is True
    assert scored["signals"]["gl_consistent"] is True


def test_weeks_payload_does_not_score(nl_export: Path):
    payload = weeks_payload(load_entries(nl_export))
    assert "results" not in payload
    assert payload["entries_in_file"] == 9


def test_invantive_purchase_lines(invantive_export: Path):
    raw = load_entries(invantive_export)
    assert {e["entry"]["entryNumber"] for e in raw["entries"]} == {"26400300", "26400387", "26200210"}
    scoped = apply_scope(raw, "auto")
    assert scoped["scope"] == "purchases"
    assert {e["entry"]["entryNumber"] for e in scoped["entries"]} == {"26400300", "26400387"}
    week37 = next(e for e in scoped["entries"] if e["entry"]["entryNumber"] == "26400387")
    assert week37["entry"]["supplierName"] == "Broekhuis Lease"
    assert week37["entry"]["amountDC"] == 1069.25
    assert week37["lines"][0]["glAccountCode"] == "4500"
    assert week37["lines"][0]["vatCode"] == "1"

    summary = review_week(scoped, "2026-W23")
    assert summary["entries"] == 1
    row = summary["results"][0]
    assert row["entryNumber"] == "26400387"
    assert row["supplier"] == "Broekhuis Lease"
    assert row["amount"] == 1069.25
    assert row["route"] in {"Auto", "AI Review", "Human Review"}
