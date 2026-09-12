"""Parse an Exact purchase export (.xlsx) into booking entries. No network."""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

WEEKDAY_NL = ["maandag", "dinsdag", "woensdag", "donderdag", "vrijdag", "zaterdag", "zondag"]
MONTH_NL = [
    "",
    "januari",
    "februari",
    "maart",
    "april",
    "mei",
    "juni",
    "juli",
    "augustus",
    "september",
    "oktober",
    "november",
    "december",
]

ALIASES = {
    "entryDate": (
        "datum",
        "boekstukdatum",
        "boekdatum",
        "entrydate",
        "date",
        "invoicedate",
        "factuurdatum",
        "transactiedatum",
        "financialyearperiod",
    ),
    "entryNumber": (
        "boekstuk",
        "boekstuknummer",
        "boekstuknr",
        "boekingnummer",
        "bookingnumber",
        "entrynumber",
        "entry",
        "entryno",
        "nummer",
        "document",
        "yourref",
    ),
    "supplierName": (
        "leverancier",
        "relatie",
        "relatienaam",
        "accountnaam",
        "accountname",
        "account",
        "supplier",
        "suppliername",
        "crediteur",
        "naam",
    ),
    "amountDC": (
        "bedrag",
        "bedragdc",
        "bedragadministratiemunteenheid",
        "bedragadministratie",
        "amountdc",
        "amount",
        "totaal",
        "amountfc",
        "bedragfc",
    ),
    "currency": ("valuta", "currency", "munteenheid"),
    "vatCode": ("btwcode", "btw", "vatcode", "vat"),
    "glAccountCode": (
        "grootboek",
        "grootboekrekeningcode",
        "grootboekrekening",
        "glaccount",
        "glaccountcode",
        "rekening",
        "rekeningcode",
    ),
    "glAccountType": (
        "grootboekrekeningsoort",
        "glaccounttype",
        "rekeningsoort",
    ),
    "journalCode": ("journaal", "dagboek", "journal", "journalcode", "dagboekcode"),
    "journalDescription": ("dagboekomschrijving", "journaldescription", "dagboeknaam"),
    "paymentCondition": (
        "betalingsconditie",
        "paymentcondition",
        "payment",
        "betaalconditie",
    ),
    "description": ("omschrijving", "description", "omschr", "tekst", "referentie"),
    "division": ("administratie", "division", "administrationcode", "divisie"),
    "lineNumber": ("regel", "linenumber", "lineno", "regelnummer"),
}

SKIP_SHEETS = {"overzicht", "overview", "readme", "parameters"}
PURCHASE_JOURNALS = {"40", "41"}
PURCHASE_JOURNAL_HINTS = ("purchase", "inkoop")
AP_GL_TYPES = {22, "22"}
COST_GL_TYPES = {110, 111, 120, 121, 122, 125, "110", "111", "120", "121", "122", "125"}
VAT_GL_CODES = {"1450", "1400", "1410", "1420"}
AP_GL_CODES = {"1300", "1600"}


def norm_header(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _alias_lookup() -> dict[str, str]:
    out = {}
    for field, names in ALIASES.items():
        for name in names:
            out[name] = field
    return out


_ALIAS = _alias_lookup()


def parse_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace("EUR", "").strip()
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text in {"", "-", "."}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_day(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(value).date().isoformat()
        except Exception:
            return None
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text[:10]):
        return text[:10]
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def iso_week_id(day: date | str) -> str:
    d = day if isinstance(day, date) else date.fromisoformat(str(day)[:10])
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def week_bounds(week: str) -> tuple[date, date]:
    parsed = parse_week(week)
    return parsed["start"], parsed["end"]


def parse_week(value: str, *, today: date | None = None) -> dict:
    """Accept 2026-W37, week 37, 37-2026, or a date that falls in that week."""
    today = today or date.today()
    text = str(value or "").strip()
    if not text:
        raise ValueError("week is empty")

    iso = re.fullmatch(r"(?:iso\s*)?(\d{4})[-.\s]?[wW](\d{1,2})", text)
    if iso:
        year, week = int(iso.group(1)), int(iso.group(2))
        return _week_from_iso(year, week)

    labeled = re.fullmatch(r"(?:week|wk)\s*(\d{1,2})(?:\s*(?:van|of)?\s*(\d{4}))?", text, re.I)
    labeled_nl = re.fullmatch(r"(\d{1,2})\s*[-/]\s*(\d{4})", text)
    if labeled:
        week = int(labeled.group(1))
        year = int(labeled.group(2) or today.year)
        return _week_from_iso(year, week)
    if labeled_nl:
        week, year = int(labeled_nl.group(1)), int(labeled_nl.group(2))
        if week > 53:
            year, week = week, year
        return _week_from_iso(year, week)

    day = parse_day(text)
    if day:
        d = date.fromisoformat(day)
        return _week_from_iso(d.isocalendar().year, d.isocalendar().week)

    raise ValueError(f"cannot parse week: {value!r}")


def _week_from_iso(year: int, week: int) -> dict:
    if week < 1 or week > 53:
        raise ValueError(f"ISO week out of range: {week}")
    start = date.fromisocalendar(year, week, 1)
    end = start + timedelta(days=6)
    week_id = f"{start.isocalendar().year}-W{start.isocalendar().week:02d}"
    return {
        "week": week_id,
        "start": start,
        "end": end,
        "label": week_label(start, end),
    }


def week_label(start: date, end: date) -> str:
    y, w, _ = start.isocalendar()
    if start.month == end.month:
        span = f"{start.day}–{end.day} {MONTH_NL[start.month]} {start.year}"
    else:
        span = (
            f"{start.day} {MONTH_NL[start.month]} – {end.day} {MONTH_NL[end.month]} {end.year}"
        )
    return f"week {w} ({span})"


def format_nl_date(day: str | date) -> str:
    d = day if isinstance(day, date) else date.fromisoformat(str(day)[:10])
    return f"{WEEKDAY_NL[d.weekday()]} {d.day} {MONTH_NL[d.month]} {d.year}"


def _sheet_rows(ws) -> list[list]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _header_score(row) -> int:
    return sum(1 for cell in row if _ALIAS.get(norm_header(cell)))


def detect_header(rows: list[list]) -> tuple[int, dict[int, str]]:
    best_i, best_score, best_map = 0, 0, {}
    scan = rows[:12] if len(rows) > 12 else rows
    for i, row in enumerate(scan):
        mapping = {}
        for col, cell in enumerate(row):
            field = _ALIAS.get(norm_header(cell))
            if field and field not in mapping.values():
                mapping[col] = field
        score = len(mapping)
        if "entryDate" in mapping.values():
            score += 2
        if score > best_score:
            best_i, best_score, best_map = i, score, mapping
    if best_score < 2 or "entryDate" not in best_map.values():
        raise ValueError(
            "Excel has no Exact transaction header (need at least a date column "
            "plus boekstuk/leverancier/bedrag)"
        )
    return best_i, best_map


def _preferred_sheets(wb) -> list:
    named = []
    rest = []
    prefer = {
        "boekingen",
        "transacties",
        "transactionlines",
        "inkoop",
        "purchases",
        "export",
        "blad1",
        "sheet1",
    }
    for ws in wb.worksheets:
        title = norm_header(ws.title)
        if title in SKIP_SHEETS:
            continue
        (named if title in prefer else rest).append(ws)
    return named + rest or list(wb.worksheets)


def _row_to_record(row, mapping: dict[int, str]) -> dict:
    rec = {field: None for field in ALIASES}
    for col, field in mapping.items():
        if col < len(row):
            rec[field] = row[col]
    rec["entryDate"] = parse_day(rec.get("entryDate"))
    rec["amountDC"] = parse_number(rec.get("amountDC"))
    rec["entryNumber"] = str(rec.get("entryNumber") or "").strip()
    rec["supplierName"] = str(rec.get("supplierName") or "").strip()
    rec["currency"] = str(rec.get("currency") or "").strip()
    rec["vatCode"] = str(rec.get("vatCode") or "").strip()
    rec["glAccountCode"] = str(rec.get("glAccountCode") or "").strip()
    rec["journalCode"] = str(rec.get("journalCode") or "").strip()
    rec["journalDescription"] = str(rec.get("journalDescription") or "").strip()
    rec["paymentCondition"] = str(rec.get("paymentCondition") or "").strip()
    rec["description"] = str(rec.get("description") or "").strip()
    rec["division"] = str(rec.get("division") or "").strip()
    rec["lineNumber"] = rec.get("lineNumber")
    rec["glAccountType"] = rec.get("glAccountType")
    return rec


def _line_role(rec: dict) -> str:
    gl = str(rec.get("glAccountCode") or "")
    kind = rec.get("glAccountType")
    line_no = str(rec.get("lineNumber") or "")
    if kind in AP_GL_TYPES or gl in AP_GL_CODES:
        return "ap"
    if gl in VAT_GL_CODES or line_no == "9999":
        return "vat"
    if kind in COST_GL_TYPES:
        return "cost"
    return "other"


def _document_amount(lines: list[dict]):
    ap = [ln["amountDC"] for ln in lines if _line_role(ln) == "ap" and ln.get("amountDC") is not None]
    if ap:
        return round(abs(sum(ap)), 2)
    costs = [ln["amountDC"] for ln in lines if _line_role(ln) == "cost" and ln.get("amountDC") is not None]
    if costs:
        return round(abs(sum(costs)), 2)
    amounts = [ln["amountDC"] for ln in lines if ln.get("amountDC") is not None]
    if not amounts:
        return None
    unique = {round(a, 2) for a in amounts}
    if len(lines) > 1 and len(unique) > 1:
        positives = [a for a in amounts if a > 0]
        return round(sum(positives), 2) if positives else round(abs(amounts[0]), 2)
    return amounts[0]


def _group_key(rec: dict, index: int) -> str:
    if rec["entryNumber"]:
        return rec["entryNumber"]
    supplier = rec["supplierName"] or "?"
    day = rec["entryDate"] or "?"
    amount = rec["amountDC"] if rec["amountDC"] is not None else "?"
    return f"{day}|{supplier}|{amount}|{index}"


def _header_from_lines(entry_number: str, lines: list[dict]) -> dict:
    first = lines[0]
    cost = next((ln for ln in lines if _line_role(ln) == "cost"), None)
    ap = next((ln for ln in lines if _line_role(ln) == "ap"), None)
    named = next((ln for ln in lines if (ln.get("supplierName") or "").strip()), first)
    primary = cost or ap or first
    descriptions = [
        ln["description"]
        for ln in (cost, ap, first)
        if ln and ln.get("description")
    ]
    return {
        "entryNumber": entry_number,
        "entryDate": first.get("entryDate"),
        "supplierName": (ap or named).get("supplierName") or "",
        "amountDC": _document_amount(lines),
        "currency": first.get("currency") or "",
        "paymentCondition": first.get("paymentCondition") or "",
        "description": descriptions[0] if descriptions else "",
        "journalCode": first.get("journalCode") or "",
        "journalDescription": first.get("journalDescription") or "",
        "journal": first.get("journalCode") or "",
        "division": first.get("division") or "",
        "vatCode": (cost or primary).get("vatCode") or "",
    }


def _as_line(rec: dict) -> dict:
    return {
        "glAccountCode": rec.get("glAccountCode") or "",
        "vatCode": rec.get("vatCode") or "",
        "amountDC": rec.get("amountDC"),
        "description": rec.get("description") or "",
        "glAccountType": rec.get("glAccountType"),
        "lineNumber": rec.get("lineNumber"),
    }


def _score_lines(recs: list[dict]) -> list[dict]:
    order = {"cost": 0, "vat": 1, "other": 2, "ap": 3}
    scored = [_as_line(r) for r in recs]
    return sorted(scored, key=lambda ln: (order.get(_line_role(ln), 2), str(ln.get("lineNumber") or "")))


def is_purchase_journal(entry: dict) -> bool:
    code = str(entry.get("journalCode") or "").strip()
    if code in PURCHASE_JOURNALS:
        return True
    desc = str(entry.get("journalDescription") or "").lower()
    return any(hint in desc for hint in PURCHASE_JOURNAL_HINTS)


REQUIRED_FIELDS = ("entryDate", "entryNumber", "supplierName", "amountDC")


def parse_report(
    entries: list[dict],
    records: list[dict],
    mapped_fields: list[str],
    mapped_headers: dict[str, str],
    sheet: str,
) -> dict:
    fallback = [
        e
        for e in entries
        if "|" in str(e["entry"].get("entryNumber") or "")
    ]
    missing_supplier = sum(1 for e in entries if not (e["entry"].get("supplierName") or "").strip())
    missing_amount = sum(1 for e in entries if e["entry"].get("amountDC") is None)
    warnings = []
    if sheet and norm_header(sheet) == "parameters":
        warnings.append("gelezen blad is Parameters — data staat op TransactionLines")
    missing_map = [f for f in REQUIRED_FIELDS if f not in mapped_fields]
    if missing_map:
        warnings.append(
            "kolom-mapping mist "
            + ", ".join(missing_map)
            + " (Invantive: Boekingnummer, Accountnaam, Bedrag Administratie Munteenheid)"
        )
    if fallback:
        warnings.append(
            f"{len(fallback)} boekstukken hebben een fallback-sleutel (datum|?|?|index): "
            "Boekingnummer is niet gemapt — niet scoren"
        )
    if entries and missing_amount > len(entries) * 0.3:
        warnings.append(
            "veel lege bedragen: map 'Bedrag Administratie Munteenheid'; "
            "som alle regels niet (debet+credit ≈ 0)"
        )
    if entries and missing_supplier > len(entries) * 0.3:
        warnings.append("veel lege leveranciers: map 'Accountnaam', niet 'Naam Abonnementhouder'")
    notes = []
    if records and entries and len(records) > len(entries) * 2:
        notes.append(
            f"{len(records)} regels → {len(entries)} boekstukken (groepering op Boekingnummer)"
        )
    ok = not any("niet scoren" in w or "mapping mist" in w for w in warnings)
    sample = [
        {
            "entryNumber": e["entry"].get("entryNumber"),
            "supplierName": e["entry"].get("supplierName"),
            "amountDC": e["entry"].get("amountDC"),
            "journalCode": e["entry"].get("journalCode"),
        }
        for e in entries[:3]
    ]
    return {
        "ok": ok,
        "sheet": sheet,
        "mappedHeaders": mapped_headers,
        "lineCount": len(records),
        "bookingCount": len(entries),
        "fallbackKeys": len(fallback),
        "missingSupplier": missing_supplier,
        "missingAmount": missing_amount,
        "warnings": warnings,
        "notes": notes,
        "sample": sample,
    }


def inspect_workbook(path: str | Path) -> dict:
    """Sheets + header mapping. Run this before asking which week."""
    path = Path(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        title = ws.title
        skip = norm_header(title) in SKIP_SHEETS
        rows = _sheet_rows(ws)
        info = {
            "name": title,
            "skip": skip,
            "rows": len(rows),
            "role": "parameters" if skip else "candidate",
        }
        if skip or not rows:
            sheets.append(info)
            continue
        try:
            header_i, mapping = detect_header(rows)
        except ValueError as exc:
            info["readable"] = False
            info["error"] = str(exc)
            sheets.append(info)
            continue
        header_row = rows[header_i]
        mapped = {
            field: str(header_row[col])
            for col, field in mapping.items()
            if col < len(header_row)
        }
        info.update(
            {
                "readable": True,
                "headerRow": header_i + 1,
                "mapped": mapped,
                "missingRequired": [f for f in REQUIRED_FIELDS if f not in mapping.values()],
            }
        )
        sheets.append(info)
    wb.close()
    chosen = next((s for s in sheets if s.get("readable") and not s.get("skip")), None)
    hints = [
        "sla Parameters over",
        "data: blad TransactionLines",
        "Boekingnummer → boekstuk (niet Factuurnummer / Rij ID)",
        "Accountnaam → leverancier (niet Naam Abonnementhouder)",
        "Bedrag Administratie Munteenheid → bedrag (niet alle regels sommeren)",
        "inkoop = dagboek 40/41",
    ]
    return {
        "source": str(path),
        "sourceName": path.name,
        "sheets": sheets,
        "chosenSheet": (chosen or {}).get("name"),
        "mapped": (chosen or {}).get("mapped") or {},
        "ready": bool(chosen) and not (chosen or {}).get("missingRequired"),
        "hints": hints,
    }


def apply_scope(bundle: dict, scope: str = "auto") -> dict:
    entries = list(bundle.get("entries") or [])
    purchases = [item for item in entries if is_purchase_journal(item["entry"])]
    used = scope
    if scope == "auto":
        used = "purchases" if purchases else "all"
    selected = purchases if used == "purchases" else entries
    out = dict(bundle)
    out["entries"] = selected
    out["scope"] = used
    out["entries_unfiltered"] = len(entries)
    return out


def load_entries(path: str | Path) -> dict:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    last_error = None
    records = []
    source_sheet = ""
    mapped_fields: list[str] = []
    mapped_headers: dict[str, str] = {}
    for ws in _preferred_sheets(wb):
        rows = _sheet_rows(ws)
        if not rows:
            continue
        try:
            header_i, mapping = detect_header(rows)
        except ValueError as exc:
            last_error = exc
            continue
        source_sheet = ws.title
        mapped_fields = sorted(mapping.values())
        mapped_headers = {
            field: str(rows[header_i][col])
            for col, field in mapping.items()
            if col < len(rows[header_i])
        }
        for offset, row in enumerate(rows[header_i + 1 :]):
            if not any(cell not in (None, "") for cell in row):
                continue
            rec = _row_to_record(row, mapping)
            if not rec["entryDate"] and rec["amountDC"] is None and not rec["supplierName"]:
                continue
            rec["_index"] = offset
            records.append(rec)
        if records:
            break
    wb.close()
    if not records:
        raise last_error or ValueError(f"no transaction rows in {path.name}")

    return bundle_from_records(
        records,
        source=str(path),
        source_name=path.name,
        sheet=source_sheet,
        mapped_fields=mapped_fields,
        mapped_headers=mapped_headers,
    )


def bundle_from_records(
    records: list[dict],
    *,
    source: str,
    source_name: str,
    sheet: str = "TransactionLines",
    mapped_fields: list[str] | None = None,
    mapped_headers: dict[str, str] | None = None,
) -> dict:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for rec in records:
        grouped[_group_key(rec, rec["_index"])].append(rec)

    entries = []
    for key, recs in grouped.items():
        recs = sorted(recs, key=lambda r: (str(r.get("lineNumber") or ""), r["_index"]))
        header = _header_from_lines(recs[0]["entryNumber"] or key, recs)
        if not header["entryDate"]:
            continue
        entries.append(
            {
                "entry": header,
                "lines": _score_lines(recs),
            }
        )
    entries.sort(key=lambda e: (e["entry"]["entryDate"] or "", e["entry"]["entryNumber"] or ""))
    divisions = [e["entry"]["division"] for e in entries if e["entry"].get("division")]
    admin = Counter(divisions).most_common(1)[0][0] if divisions else ""
    mapped_fields = mapped_fields or sorted(
        {k for rec in records for k in rec if rec.get(k) not in (None, "")}
    )
    mapped_headers = mapped_headers or {field: field for field in mapped_fields}
    parse = parse_report(entries, records, mapped_fields, mapped_headers, sheet)
    return {
        "source": source,
        "sourceName": source_name,
        "sheet": sheet,
        "administrationCode": admin,
        "entries": entries,
        "rowCount": len(records),
        "mappedFields": mapped_fields,
        "mappedHeaders": mapped_headers,
        "parse": parse,
    }


def list_weeks(bundle: dict) -> list[dict]:
    buckets: dict[str, list] = defaultdict(list)
    for item in bundle["entries"]:
        day = item["entry"].get("entryDate")
        if not day:
            continue
        buckets[iso_week_id(day)].append(item)
    weeks = []
    for week_id, items in sorted(buckets.items()):
        start, end = week_bounds(week_id)
        weeks.append(
            {
                "week": week_id,
                "start": start.isoformat(),
                "end": end.isoformat(),
                "label": week_label(start, end),
                "entries": len(items),
            }
        )
    return weeks


def entries_in_week(bundle: dict, week: str) -> tuple[dict, list[dict]]:
    meta = parse_week(week)
    start, end = meta["start"], meta["end"]
    matched = []
    for item in bundle["entries"]:
        day = item["entry"].get("entryDate")
        if not day:
            continue
        d = date.fromisoformat(day)
        if start <= d <= end:
            matched.append(item)
    return meta, matched


def _supplier_key(name: str) -> str:
    return re.sub(r"\s+", " ", (name or "").strip().lower())


def build_cases(bundle: dict, week_entries: list[dict]) -> list[dict]:
    """History comes from other rows in the same uploaded workbook."""
    by_supplier: dict[str, list[dict]] = defaultdict(list)
    for item in bundle["entries"]:
        by_supplier[_supplier_key(item["entry"].get("supplierName"))].append(item)

    cases = []
    for item in week_entries:
        entry = item["entry"]
        current_no = str(entry.get("entryNumber") or "")
        current_day = entry.get("entryDate") or ""
        peers = [
            p
            for p in by_supplier[_supplier_key(entry.get("supplierName"))]
            if str(p["entry"].get("entryNumber") or "") != current_no
            or (not current_no and p is not item)
        ]
        peers.sort(key=lambda p: (p["entry"].get("entryDate") or "", p["entry"].get("entryNumber") or ""))
        prior = [p for p in peers if (p["entry"].get("entryDate") or "") <= current_day]
        history_src = prior if prior else peers
        history = [p["entry"] for p in history_src]
        recent = history_src[-3:]
        history_lines = [line for p in recent for line in p.get("lines") or []]
        supplier = (entry.get("supplierName") or "").strip()
        cases.append(
            {
                "entry": entry,
                "lines": item.get("lines") or [],
                "history": history,
                "history_lines": history_lines,
                "accounts_count": 1 if supplier else 0,
            }
        )
    return cases
