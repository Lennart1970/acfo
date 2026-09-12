#!/usr/bin/env python3
"""Write a specialist Excel dagoverzicht from a LedgerBotje purchase-search payload.

Sandbox-safe: no network. Copilot Studio GitHub-harness skills run this locally
after the agent has already fetched Exact via the ledgerbotje *tool*.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

WEEKDAY_NL = ["ma", "di", "wo", "do", "vr", "za", "zo"]
MONTH_NL = [
    "", "januari", "februari", "maart", "april", "mei", "juni",
    "juli", "augustus", "september", "oktober", "november", "december",
]

COLUMNS = [
    "Datum",
    "Boekstuk",
    "Leverancier",
    "Bedrag",
    "Valuta",
    "Btw-code",
    "Journaal",
    "Betalingsconditie",
    "Omschrijving",
    "Opvallend",
]

SCORED_COLUMNS = COLUMNS + ["Route", "Confidence", "Signalen"]

ROUTE_FILL = {
    "Human Review": "F4B183",
    "AI Review": "FFE699",
    "Auto": "C6EFCE",
}


def format_nl_date(day: str) -> str:
    d = date.fromisoformat(str(day)[:10])
    return f"{WEEKDAY_NL[d.weekday()]} {d.day} {MONTH_NL[d.month]} {d.year}"


def load_payload(raw) -> dict:
    if isinstance(raw, str):
        raw = json.loads(raw)
    if not isinstance(raw, dict):
        raise ValueError("payload must be a JSON object")
    # MCP wrapper: {result: {content: [{text: "...json..."}]}}
    if "results" not in raw and "result" in raw:
        content = (raw.get("result") or {}).get("content") or []
        texts = [c.get("text") for c in content if isinstance(c, dict) and c.get("type", "text") == "text"]
        if texts:
            try:
                raw = json.loads(texts[0])
            except json.JSONDecodeError:
                pass
    return raw


def extract_results(payload: dict) -> tuple[list, str]:
    payload = load_payload(payload)
    if isinstance(payload.get("results"), list):
        admin = str((payload.get("activeContext") or {}).get("administrationCode") or "")
        return payload["results"], admin
    if isinstance(payload.get("matched"), list):
        return payload["matched"], str(payload.get("administrationCode") or "")
    if isinstance(payload.get("entries"), list):
        return payload["entries"], str(payload.get("administrationCode") or "")
    raise ValueError("JSON has no results/matched/entries array")


def flags_for(row: dict) -> str:
    notes = []
    supplier = str(row.get("supplierName") or row.get("supplier") or "").strip()
    desc = str(row.get("description") or "").strip()
    amount = row.get("amountDC")
    if amount is None:
        amount = row.get("amount")
    if not supplier:
        notes.append("geen leverancier")
    if amount in (None, 0, 0.0, "0"):
        notes.append("bedrag 0 of leeg")
    if not desc:
        notes.append("geen omschrijving")
    return "; ".join(notes)


def row_date(row: dict) -> str:
    return str(row.get("entryDate") or row.get("date") or "")[:10]


def to_rows(entries: list, day: str) -> list[dict]:
    matched = [e for e in entries if row_date(e) == day]
    out = []
    for e in matched:
        amount = e.get("amountDC")
        if amount is None:
            amount = e.get("amount")
        out.append({
            "Datum": row_date(e),
            "Boekstuk": str(e.get("entryNumber") or ""),
            "Leverancier": str(e.get("supplierName") or e.get("supplier") or "").strip(),
            "Bedrag": amount,
            "Valuta": e.get("currency") or "",
            "Btw-code": e.get("vatCode") or "",
            "Journaal": e.get("journal") or e.get("journalCode") or "",
            "Betalingsconditie": e.get("paymentCondition") or "",
            "Omschrijving": e.get("description") or "",
            "Opvallend": flags_for(e),
        })
    return out


def write_xlsx(path: Path, day: str, admin: str, searched: int, rows: list[dict], *, scored: bool = False) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise SystemExit("openpyxl is required in the Copilot sandbox") from exc

    wb = Workbook()
    overview = wb.active
    overview.title = "Overzicht"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    overview["A1"] = "aCFO booking check" if scored else "aCFO dagoverzicht inkoop"
    overview["A1"].font = Font(bold=True, size=14)
    kind = "Booking Confidence (Python in sandbox — Exact niet gewijzigd)" if scored else "dagoverzicht — geen Booking Confidence-score"
    counts = ""
    if scored:
        from collections import Counter
        c = Counter(r.get("Route") or "" for r in rows)
        counts = f"Auto {c.get('Auto', 0)} / AI Review {c.get('AI Review', 0)} / Human Review {c.get('Human Review', 0)}"
    meta = [
        ("Datum", format_nl_date(day), day),
        ("Administratie", admin or "(actieve Exact-administratie)", ""),
        ("Opgehaald (recent, max 50)", searched, ""),
        ("Boekingen op datum", len(rows), ""),
        ("Routes", counts or "—", ""),
        ("Exact gewijzigd", "nee", ""),
        ("Dit is", kind, ""),
    ]
    overview["A3"] = "Veld"
    overview["B3"] = "Waarde"
    for col in ("A3", "B3"):
        overview[col].font = header_font
        overview[col].fill = header_fill
    for i, (label, value, _) in enumerate(meta, start=4):
        overview[f"A{i}"] = label
        overview[f"B{i}"] = value

    columns = SCORED_COLUMNS if scored else COLUMNS
    bookings = wb.create_sheet("Boekingen")
    for col, name in enumerate(columns, start=1):
        cell = bookings.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    for r_i, row in enumerate(rows, start=2):
        for c_i, name in enumerate(columns, start=1):
            cell = bookings.cell(r_i, c_i, row.get(name))
            cell.border = thin
            if name == "Opvallend" and row.get(name):
                cell.fill = warn_fill
            if name == "Route" and row.get(name) in ROUTE_FILL:
                cell.fill = PatternFill("solid", fgColor=ROUTE_FILL[row[name]])
            if name == "Bedrag" and isinstance(row.get(name), (int, float)):
                cell.number_format = "#,##0.00"
            if name == "Confidence" and isinstance(row.get(name), (int, float)):
                cell.number_format = "0%"
        bookings.cell(r_i, 1).number_format = "YYYY-MM-DD"
    widths = {
        "A": 14, "B": 14, "C": 32, "D": 12, "E": 10,
        "F": 12, "G": 14, "H": 20, "I": 40, "J": 28,
        "K": 16, "L": 12, "M": 36,
    }
    for col, width in widths.items():
        bookings.column_dimensions[col].width = width
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 55

    bookings.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    bookings.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build(payload, day: str, admin: str | None = None) -> dict:
    entries, payload_admin = extract_results(payload)
    rows = to_rows(entries, day)
    return {
        "date": day,
        "administrationCode": admin or payload_admin,
        "searched": len(entries),
        "entries": len(rows),
        "rows": rows,
    }


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="Write BookingCheck-YYYY-MM-DD.xlsx from purchase-search JSON")
    p.add_argument("--date", required=True, help="YYYY-MM-DD")
    p.add_argument("--admin", default="", help="Exact administration code")
    p.add_argument("--input", required=True, help="Path to LedgerBotje search JSON")
    p.add_argument("--output", default="", help="xlsx path (default BookingCheck-DATE.xlsx)")
    p.add_argument("--json", action="store_true", help="also print summary JSON to stdout")
    return p.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    summary = build(payload, args.date, args.admin or None)
    out = Path(args.output) if args.output else Path(f"BookingCheck-{args.date}.xlsx")
    write_xlsx(out, summary["date"], summary["administrationCode"], summary["searched"], summary["rows"])
    if args.json:
        dump = {k: v for k, v in summary.items() if k != "rows"}
        dump["output"] = str(out)
        dump["boekstukken"] = [r["Boekstuk"] for r in summary["rows"]]
        json.dump(dump, sys.stdout, ensure_ascii=False, indent=2)
        sys.stdout.write("\n")
    else:
        print(f"wrote {out} ({summary['entries']} of {summary['searched']} on {args.date})", file=sys.stderr)
    return 0 if summary["entries"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
