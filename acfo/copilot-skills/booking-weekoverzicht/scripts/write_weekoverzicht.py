"""Write a weekly booking-review Excel from scored purchase rows. No network."""
from __future__ import annotations

from pathlib import Path

from excel_transactions import format_nl_date, week_label

COLUMNS = [
    "Datum",
    "Boekstuk",
    "Leverancier",
    "Bedrag",
    "Valuta",
    "Btw-code",
    "Journaal",
    "Betalingsconditie",
    "Omschrijving",
    "Opvallend",
    "Route",
    "Confidence",
    "Signalen",
    "Uitkomst",
]

ROUTE_FILL = {
    "Human Review": "F4B183",
    "AI Review": "FFE699",
    "Auto": "C6EFCE",
}

SIGNAL_NL = {
    "known_supplier": "bekende leverancier",
    "booked_before": "eerder geboekt",
    "gl_consistent": "grootboek",
    "vat_consistent": "btw",
    "amount_in_range": "bedrag",
    "cadence_consistent": "timing",
    "description_similar": "omschrijving",
    "currency_payment_normal": "valuta/betaling",
    "no_unusual_change": "geen bijzondere afwijking",
    "signals_agree": "signalen eens",
}


def flags_for(row: dict) -> str:
    notes = []
    supplier = str(row.get("supplierName") or row.get("supplier") or "").strip()
    desc = str(row.get("description") or "").strip()
    amount = row.get("amountDC")
    if amount is None:
        amount = row.get("amount")
    if not supplier:
        notes.append("geen leverancier")
    if amount in (None, 0, 0.0, "0"):
        notes.append("bedrag 0 of leeg")
    if not desc:
        notes.append("geen omschrijving")
    return "; ".join(notes)


def failed_signals(signals: dict) -> str:
    failed = [SIGNAL_NL.get(k, k) for k, v in signals.items() if k != "signals_agree" and not v]
    return ", ".join(failed)


def outcome_text(scored: dict) -> str:
    route = scored["route"]
    confidence = scored["confidence"]
    failed = failed_signals(scored["signals"])
    if route == "Auto":
        text = f"Auto — vertrouwde boeking ({confidence:.0%})"
    elif route == "AI Review":
        text = f"AI Review — twijfel, extra check ({confidence:.0%})"
    else:
        text = f"Human Review — handmatig beoordelen ({confidence:.0%})"
    if failed:
        text += f". Afwijkend: {failed}"
    return text


def scored_row(scored: dict) -> dict:
    entry = scored["entry"]
    lines = scored.get("lines") or []
    ev = scored["evidence"]
    vat = ev["current"].get("vat") or (lines[0].get("vatCode") if lines else "")
    gl = ev["current"].get("gl")
    return {
        "Datum": str(entry.get("entryDate") or "")[:10],
        "Boekstuk": str(entry.get("entryNumber") or ""),
        "Leverancier": ev.get("supplier") or "",
        "Bedrag": entry.get("amountDC"),
        "Valuta": entry.get("currency") or "",
        "Btw-code": vat or "",
        "Journaal": entry.get("journalCode") or entry.get("journal") or gl or "",
        "Betalingsconditie": entry.get("paymentCondition") or "",
        "Omschrijving": entry.get("description") or "",
        "Opvallend": flags_for(entry),
        "Route": scored["route"],
        "Confidence": scored["confidence"],
        "Signalen": failed_signals(scored["signals"]),
        "Uitkomst": outcome_text(scored),
    }


def write_xlsx(path: Path, summary: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = Path(path)
    rows = summary["rows"]
    wb = Workbook()
    overview = wb.active
    overview.title = "Overzicht"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    warn_fill = PatternFill("solid", fgColor="FFF2CC")

    overview["A1"] = "aCFO weekoverzicht inkoop"
    overview["A1"].font = Font(bold=True, size=14)

    counts = summary.get("counts") or {}
    routes = (
        f"Auto {counts.get('Auto', 0)} / "
        f"AI Review {counts.get('AI Review', 0)} / "
        f"Human Review {counts.get('Human Review', 0)}"
    )
    start = summary.get("start") or ""
    end = summary.get("end") or ""
    period = ""
    if start and end:
        period = f"{format_nl_date(start)} – {format_nl_date(end)}"
    label = summary.get("label") or (
        week_label(
            __import__("datetime").date.fromisoformat(start),
            __import__("datetime").date.fromisoformat(end),
        )
        if start and end
        else summary.get("week") or ""
    )
    meta = [
        ("Week", label),
        ("Periode", period),
        ("Administratie", summary.get("administrationCode") or "(uit Excel)"),
        ("Bronbestand", summary.get("sourceName") or ""),
        ("Boekingen in week", len(rows)),
        ("Routes", routes),
        ("Exact gewijzigd", "nee"),
        (
            "Dit is",
            "Booking Confidence v2 — Excel-upload, Python in sandbox, Exact niet geraadpleegd",
        ),
    ]
    overview["A3"] = "Veld"
    overview["B3"] = "Waarde"
    for col in ("A3", "B3"):
        overview[col].font = header_font
        overview[col].fill = header_fill
    for i, (label_cell, value) in enumerate(meta, start=4):
        overview[f"A{i}"] = label_cell
        overview[f"B{i}"] = value
    overview.column_dimensions["A"].width = 28
    overview.column_dimensions["B"].width = 78

    bookings = wb.create_sheet("Boekingen")
    for col, name in enumerate(COLUMNS, start=1):
        cell = bookings.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    for r_i, row in enumerate(rows, start=2):
        for c_i, name in enumerate(COLUMNS, start=1):
            cell = bookings.cell(r_i, c_i, row.get(name))
            cell.border = thin
            if name == "Opvallend" and row.get(name):
                cell.fill = warn_fill
            if name == "Route" and row.get(name) in ROUTE_FILL:
                cell.fill = PatternFill("solid", fgColor=ROUTE_FILL[row[name]])
            if name == "Bedrag" and isinstance(row.get(name), (int, float)):
                cell.number_format = "#,##0.00"
            if name == "Confidence" and isinstance(row.get(name), (int, float)):
                cell.number_format = "0%"
            if name == "Uitkomst":
                cell.alignment = Alignment(wrap_text=True)
        bookings.cell(r_i, 1).number_format = "YYYY-MM-DD"
    widths = {
        "A": 14,
        "B": 14,
        "C": 32,
        "D": 12,
        "E": 10,
        "F": 12,
        "G": 14,
        "H": 20,
        "I": 36,
        "J": 24,
        "K": 16,
        "L": 12,
        "M": 36,
        "N": 64,
    }
    for col, width in widths.items():
        bookings.column_dimensions[col].width = width
    bookings.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"
    bookings.freeze_panes = "A2"

    outcomes = wb.create_sheet("Uitkomsten")
    outcomes["A1"] = "Uitkomst per transactie"
    outcomes["A1"].font = Font(bold=True, size=14)
    out_cols = [
        "Boekstuk",
        "Datum",
        "Leverancier",
        "Bedrag",
        "Route",
        "Confidence",
        "Uitkomst",
    ]
    for col, name in enumerate(out_cols, start=1):
        cell = outcomes.cell(3, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for r_i, row in enumerate(rows, start=4):
        values = [
            row.get("Boekstuk"),
            row.get("Datum"),
            row.get("Leverancier"),
            row.get("Bedrag"),
            row.get("Route"),
            row.get("Confidence"),
            row.get("Uitkomst"),
        ]
        for c_i, value in enumerate(values, start=1):
            cell = outcomes.cell(r_i, c_i, value)
            cell.border = thin
            if out_cols[c_i - 1] == "Route" and value in ROUTE_FILL:
                cell.fill = PatternFill("solid", fgColor=ROUTE_FILL[value])
            if out_cols[c_i - 1] == "Bedrag" and isinstance(value, (int, float)):
                cell.number_format = "#,##0.00"
            if out_cols[c_i - 1] == "Confidence" and isinstance(value, (int, float)):
                cell.number_format = "0%"
            if out_cols[c_i - 1] == "Uitkomst":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFG", (14, 14, 32, 12, 16, 12, 72)):
        outcomes.column_dimensions[col].width = width
    outcomes.row_dimensions[3].height = 18
    outcomes.freeze_panes = "A4"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
